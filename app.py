import streamlit as st
import requests
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import re
from bs4 import BeautifulSoup
from urllib.parse import urljoin

st.set_page_config(page_title="表演藝術演出節目統計系統", page_icon="🎭", layout="wide")

st.markdown("""
<style>
    .main-title { font-size: 26px; font-weight: bold; color: #1E3A8A; margin-bottom: 5px; }
    .sub-title { font-size: 14px; color: #4B5563; margin-bottom: 20px; }
    .stTabs [data-baseweb="tab-list"] { gap: 15px; }
    .stTabs [data-baseweb="tab"] { font-size: 16px; font-weight: bold; }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">🎭 表演藝術與展演售票節目統計系統</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">支援 UDN 官網直連、文化部全國資料庫、ACCUPASS、智慧藝術節識別與雙分頁 Excel 匯出</div>', unsafe_allow_html=True)

PLATFORM_OPTIONS = {
    "🎪 udn 售票網 (官網直連)": "UDN",
    "🌐 全部整合平台 (全台聯播)": "ALL",
    "🎟️ OPENTIX (兩廳院文化生活)": "OPENTIX",
    "🎫 ACCUPASS (活動通專屬)": "ACCUPASS",
    "🎸 KKTIX": "KKTIX",
    "🎺 寬宏售票": "KHAM",
    "🎻 年代售票": "ERA"
}

CATEGORY_MAP_CULTURE = {
    "音樂": "1",
    "戲劇": "2",
    "舞蹈": "3",
    "戲曲": "17",
    "親子": "7",
    "其他": "OTHER"
}

CATEGORY_MAP_UDN = {
    "戲劇": "116",
    "音樂": "77",
    "舞蹈": "100",
    "親子": "129"
}

EXHIBITION_EXCLUDE = [
    "特展", "展覽", "文物展", "光影展", "恐龍展", "珠寶展", "書展", "動漫展",
    "快閃店", "沉浸展", "藝術展", "公仔展", "沉浸式藝術", "世界巡迴展", "講座", "研討會",
    "法會", "道場", "祈福法會", "祈安", "招生講座"
]

ACCUPASS_EXCLUDE_KEYWORDS = [
    "工作坊", "手作", "市集", "體驗課", "教學", "培訓", "說明會", "品酒",
    "創業", "行銷", "分享會", "讀書會", "烘焙", "證照", "瑜珈", "運動會", "線上直播", "論壇"
]

BUILTIN_FESTIVALS = [
    "TIFA", "NTT-TIFA", "KSAF", "TTTF", "Fringe", "藝穗節", "台北藝術節", "臺北藝術節",
    "新點子實驗場", "衛武營小時光", "夏日童樂會", "國際劇場藝術節", "臺灣戲曲藝術節",
    "大稻埕戲苑傳統戲曲季", "新北兒童藝術節", "夏至藝術節", "秋天藝術節", "高雄春天藝術節",
    "台南藝術節", "臺南藝術節", "花蓮太平洋左岸藝術節", "嘉義市瘋科技藝術節"
]

WELL_KNOWN_GROUPS = [
    "果陀劇場", "綠光劇團", "明華園", "唐美雲歌仔戲團", "國光劇團", "當代傳奇劇場",
    "故事工廠", "春河劇團", "台南人劇團", "全民大劇團", "相聲瓦舍", "無垢舞蹈劇場",
    "雲門舞集", "布拉瑞揚舞團", "國家交響樂團", "台北市立交響樂團", "高雄市交響樂團",
    "台北民族舞團", "薪傳歌仔戲劇團", "一心戲劇團", "秀琴歌劇團", "紙風車劇團",
    "如果兒童劇團", "偶偶偶劇團", "鞋子兒童實驗劇團", "卡米地", "薩泰爾", "漫才少爺",
    "FOCA福爾摩沙馬戲團", "驚喜製造", "衛武營", "國家兩廳院", "臺中國家歌劇院"
]

def extract_festival(title, desc="", user_keywords=[]):
    combined = f"{title} {desc[:120]}"
    for kw in user_keywords:
        kw = kw.strip()
        if kw and (kw.lower() in combined.lower()):
            return kw
    for fest in BUILTIN_FESTIVALS:
        pattern = re.compile(rf"(?:20\d\d\s*)?{re.escape(fest)}", re.IGNORECASE)
        m = pattern.search(title)
        if m:
            return m.group(0).strip()
        if fest.lower() in combined.lower():
            return fest
    symbol_match = re.search(r"^[【\[［]([^】\]］]+)[】\]］]", title)
    if symbol_match:
        content = symbol_match.group(1).strip()
        if any(term in content for term in ["節", "季", "祭", "週", "Festival", "系列", "展演"]):
            return content
    dash_match = re.search(r"^([^－—\-]+)[－—\-]", title)
    if dash_match:
        content = dash_match.group(1).strip()
        if any(term in content for term in ["節", "季", "祭", "週", "Festival", "小時光", "童樂會"]):
            if len(content) <= 25:
                return content
    suffix_pattern = r"(?:20\d\d\s*)?[\u4e00-\u9fa5A-Za-z0-9]{2,12}(?:藝術節|藝術季|戲劇節|舞蹈節|音樂節|音樂祭|文化季|藝穗節|生活節|劇節|舞季|戲曲季|戲曲節|兒童節|Festival)"
    m_suffix = re.search(suffix_pattern, title, re.IGNORECASE)
    if m_suffix:
        found = m_suffix.group(0).strip()
        if len(found) >= 4:
            return found
    return "單檔演出 (一般檔期)"

def extract_performer(show_title, master_unit=[], sub_unit=[], other_unit=[], desc=""):
    match = re.search(r"[【\[［]([^】\]］]+)[】\]］]", show_title)
    if match:
        candidate = match.group(1).strip()
        if any(term in candidate for term in ["劇團", "舞團", "樂團", "合唱團", "工作室", "劇會", "劇坊", "劇社", "製作", "馬戲"]):
            return candidate

    match_book = re.search(r"^([^\d《]+)《", show_title)
    if match_book:
        candidate = match_book.group(1).strip()
        if len(candidate) <= 15 and any(term in candidate for term in ["團", "社", "坊", "會", "卡米地"]):
            return candidate

    combined_text = f"{show_title} {desc[:80]}"
    for group in WELL_KNOWN_GROUPS:
        if group in combined_text:
            return group

    if other_unit:
        if isinstance(other_unit, list):
            clean = [u.strip() for u in other_unit if u.strip()]
            if clean: return "、".join(clean)
        elif isinstance(other_unit, str) and other_unit.strip():
            return other_unit.strip()

    if master_unit:
        m_str = "、".join(master_unit) if isinstance(master_unit, list) else str(master_unit)
        if not any(gov in m_str for gov in ["文化局", "文化部", "公所", "市政府", "縣政府", "UDN"]):
            return m_str.strip()

    return "詳見售票頁"

def identify_platform_deep(web_sales_url, desc="", title="", comment=""):
    combined = f"{web_sales_url} {desc} {title} {comment}".lower()
    if any(k in combined for k in ["opentix.life", "opentix", "兩廳院售票", "兩廳院文化生活"]):
        return "OPENTIX"
    elif any(k in combined for k in ["kktix.cc", "kktix.com", "kktix"]):
        return "KKTIX"
    elif any(k in combined for k in ["tickets.udnfunlife.com", "udn售票", "udn.com", "聯合數位文創", "udn 售票"]):
        return "UDN"
    elif any(k in combined for k in ["kham.com.tw", "寬宏售票", "寬宏藝術"]):
        return "KHAM"
    elif any(k in combined for k in ["ticket.com.tw", "年代售票"]):
        return "ERA"
    elif "accupass.com" in combined or "活動通" in combined:
        return "ACCUPASS"
    return "其他售票通路"

# 側邊欄控制項
with st.sidebar:
    st.header("⚙️ 統計篩選條件")
    selected_platform_label = st.selectbox("1. 選擇售票平台：", options=list(PLATFORM_OPTIONS.keys()), index=0)
    target_platform = PLATFORM_OPTIONS[selected_platform_label]

    selected_cats = st.multiselect(
        "2. 選擇節目種類：",
        options=["戲劇", "音樂", "舞蹈", "親子", "戲曲", "其他"],
        default=["戲劇", "音樂", "舞蹈", "親子"]
    )

    today = datetime.date.today()
    default_end = today + datetime.timedelta(days=120)

    st.write("3. 演出日期區間：")
    c1, c2 = st.columns(2)
    with c1:
        start_date = st.date_input("開始日期", today)
    with c2:
        end_date = st.date_input("結束日期", default_end)

    st.write("4. 補充藝術節/季關鍵字 (選填)：")
    custom_festivals_input = st.text_input("以逗號分隔，如：Sound Wave, 某某祭", "")
    custom_fest_list = [k.strip() for k in custom_festivals_input.split(",") if k.strip()]

    start_btn = st.button("🚀 開始擷取並產生統計", type="primary", use_container_width=True)

# UDN 專屬字串解析
def parse_udn_card_text(text):
    date_match = re.search(r"(\d{4}[./-]\d{1,2}[./-]\d{1,2}(?:\s*~\s*\d{4}[./-]\d{1,2}[./-]\d{1,2})?)", text)
    date_str = date_match.group(1).strip() if date_match else "即期 (詳見售票頁)"
    title = text
    venue = "詳見售票頁"
    if date_match:
        idx_start, idx_end = date_match.span()
        title = text[:idx_start].strip()
        after_date = text[idx_end:].strip()
        venue_match = re.split(r"NT\s*\$|NTD|\$", after_date)
        venue = venue_match[0].replace("...more", "").strip() if venue_match else after_date
    return title, date_str, venue

# UDN 爬蟲
def fetch_udn_shows_direct(selected_categories, pbar, status_box):
    udn_records = []
    base_url = "https://tickets.udnfunlife.com"
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"}
    target_routes = []
    for cat_name in ["戲劇", "音樂", "舞蹈", "親子"]:
        if cat_name in selected_categories:
            cat_id = CATEGORY_MAP_UDN[cat_name]
            url = f"https://tickets.udnfunlife.com/application/UTK01/UTK0101_03.aspx?Category={cat_id}&kdid=cateList"
            target_routes.append((url, cat_name))

    total = len(target_routes)
    for idx, (url, cat_name) in enumerate(target_routes):
        status_box.info(f"正在擷取 UDN【{cat_name}】專區在售節目...")
        try:
            r = requests.get(url, headers=headers, timeout=12)
            if r.status_code == 200:
                soup = BeautifulSoup(r.text, "lxml")
                a_tags = soup.find_all("a", href=re.compile(r"PRODUCT_ID="))
                seen_links = set()
                for a in a_tags:
                    href = a.get("href", "")
                    full_link = urljoin(base_url, href)
                    if full_link in seen_links: continue
                    seen_links.add(full_link)
                    raw_text = a.get_text(" ", strip=True)
                    if not raw_text or len(raw_text) < 4: continue
                    title, date_str, venue = parse_udn_card_text(raw_text)
                    if not title or len(title) < 2: continue

                    udn_records.append({
                        "title": title,
                        "category": cat_name,
                        "date_str": date_str,
                        "venue": venue if venue else "詳見售票頁",
                        "link": full_link
                    })
        except Exception:
            pass
        pbar.progress(int((idx + 1) / max(total, 1) * 80))
    return udn_records

# ACCUPASS 直連
def fetch_accupass_shows(start_date, end_date):
    accupass_records = []
    headers = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)", "Referer": "https://www.accupass.com/"}
    for kw in ["戲劇", "音樂", "演出", "脫口秀", "讀劇", "喜劇", "漫才", "沉浸式", "舞蹈"]:
        try:
            url = f"https://api.accupass.com/v3/search/events?keyword={kw}&page=0&size=25"
            r = requests.get(url, headers=headers, timeout=8)
            if r.status_code == 200:
                events = r.json().get("events", [])
                for ev in events:
                    title = ev.get("title", "").strip()
                    summary = ev.get("summary", "")
                    if any(ex in f"{title} {summary}" for ex in ACCUPASS_EXCLUDE_KEYWORDS):
                        continue
                    event_id = ev.get("eventId", "")
                    org_name = ev.get("orgName", "ACCUPASS 主辦單位")
                    event_url = f"https://www.accupass.com/event/{event_id}" if event_id else "https://www.accupass.com"
                    start_time = ev.get("eventStartDateTime", "")
                    address = ev.get("address", "") or ev.get("location", "詳見活動頁面")
                    accupass_records.append({
                        "title": title,
                        "category": "其他",
                        "masterUnit": [org_name],
                        "link": event_url,
                        "date_str": start_time if start_time else "詳見售票頁",
                        "venue": address
                    })
        except Exception:
            pass
    return accupass_records

# 執行流程
if start_btn:
    if not selected_cats:
        st.warning("⚠️ 請至少勾選一種節目種類！")
    elif start_date > end_date:
        st.error("❌ 開始日期不能晚於結束日期！")
    else:
        status_box = st.empty()
        pbar = st.progress(0)
        detail_records = []

        # 分流 1：UDN 官網直連
        if target_platform == "UDN":
            status_box.info("正在連線 UDN 售票網官網即時擷取演出節目...")
            shows = fetch_udn_shows_direct(selected_cats, pbar, status_box)

            for s in shows:
                title = s["title"]
                cat = s["category"]
                link = s["link"]
                date_str = s["date_str"]
                venue = s["venue"]
                fest = extract_festival(title, "", custom_fest_list)
                performer = extract_performer(title)

                m_date = re.search(r"(\d{4}[./-]\d{1,2}[./-]\d{1,2})", date_str)
                first_date = m_date.group(1).replace(".", "-").replace("/", "-") if m_date else "即期"

                detail_records.append({
                    "售票通路": "UDN 售票網",
                    "所屬藝術節/系列": fest,
                    "節目分類": cat,
                    "演出單位": performer,
                    "節目名稱": title,
                    "演出場次時間": date_str,
                    "演出地點": venue,
                    "售票狀況/連結": link
                })

        # 分流 2：ACCUPASS 直連
        elif target_platform == "ACCUPASS":
            status_box.info("正在連線 ACCUPASS 專屬活動通道並進行表演除噪...")
            shows = fetch_accupass_shows(start_date, end_date)
            pbar.progress(70)

            for s in shows:
                title = s["title"]
                cat = "其他"
                fest = extract_festival(title, "", custom_fest_list)
                performer = extract_performer(title, master_unit=s.get("masterUnit", []))

                detail_records.append({
                    "售票通路": "ACCUPASS",
                    "所屬藝術節/系列": fest,
                    "節目分類": cat,
                    "演出單位": performer,
                    "節目名稱": title,
                    "演出場次時間": s["date_str"],
                    "演出地點": s["venue"],
                    "售票狀況/連結": s["link"]
                })

        # 分流 3：文化部全台統整大盤 (OPENTIX / 寬宏 / 年代 / KKTIX / 全部)
        else:
            status_box.info(f"正在連線全國藝文活動資料庫篩選【{selected_platform_label}】節目...")
            API_URL = "https://cloud.culture.tw/frontsite/trans/SearchShowAction.do?method=doFindTypeJ"
            raw_data = []

            fetch_categories = set()
            for cat in selected_cats:
                if cat == "其他":
                    fetch_categories.update(["1", "2", "3", "5", "6", "17", "7"])
                else:
                    fetch_categories.add(CATEGORY_MAP_CULTURE[cat])

            fetch_list = list(fetch_categories)
            total = len(fetch_list)
            for i, code in enumerate(fetch_list):
                try:
                    r = requests.get(API_URL, params={"category": code}, timeout=15)
                    if r.status_code == 200:
                        raw_data.extend(r.json())
                except Exception:
                    pass
                pbar.progress(int((i + 1) / max(total, 1) * 60))

            seen_ids = set()
            for s in raw_data:
                uid = s.get("UID") or s.get("title", "")
                if uid in seen_ids: continue
                seen_ids.add(uid)

                title = s.get("title", "").strip()
                if any(ex in title for ex in EXHIBITION_EXCLUDE): continue

                desc = s.get("descriptionFilterHtml", "") or s.get("comment", "")
                web_sales = s.get("webSales", "").strip()
                comment_text = s.get("comment", "")

                # 平台深度篩選
                detected_platform = identify_platform_deep(web_sales, desc, title, comment_text)
                if target_platform != "ALL" and detected_platform != target_platform:
                    continue

                cat_code = s.get("category", "")
                cat_name = "其他"
                for cn, cc in CATEGORY_MAP_CULTURE.items():
                    if cc == cat_code:
                        cat_name = cn
                        break

                if cat_name not in selected_cats:
                    continue

                performer = extract_performer(title, master_unit=s.get("masterUnit", []), sub_unit=s.get("subUnit", []), other_unit=s.get("otherUnit", []), desc=desc)
                festival = extract_festival(title, desc, custom_fest_list)

                shows_info = s.get("showInfo", [])
                if not shows_info:
                    shows_info = [{"time": "詳見售票頁", "locationName": "詳見售票頁", "location": ""}]

                for si in shows_info:
                    ptime = si.get("time", "").strip()
                    loc_name = si.get("locationName", "").strip()
                    loc_detail = si.get("location", "").strip()
                    floc = f"{loc_name} ({loc_detail})" if (loc_detail and loc_detail != loc_name) else loc_name

                    if len(ptime) >= 10:
                        try:
                            sdt = datetime.datetime.strptime(ptime[:10].replace("/", "-"), "%Y-%m-%d").date()
                            if not (start_date <= sdt <= end_date):
                                continue
                        except Exception:
                            pass

                    detail_records.append({
                        "售票通路": detected_platform if target_platform != "ALL" else "全台售票通路",
                        "所屬藝術節/系列": festival,
                        "節目分類": cat_name,
                        "演出單位": performer,
                        "節目名稱": title,
                        "演出場次時間": ptime if ptime else "詳見售票頁",
                        "演出地點": floc if floc else "詳見售票頁",
                        "售票狀況/連結": web_sales if web_sales else "請洽售票平台/主辦單位"
                    })

        pbar.progress(85)

        if not detail_records:
            status_box.empty()
            pbar.empty()
            st.warning(f"💡 在所選的【{selected_platform_label}】平台與條件下，未找到符合條件之節目。建議放寬日期區間或節目種類再次查詢。")
        else:
            df_details = pd.DataFrame(detail_records)

            summary_list = []
            grouped = df_details.groupby("節目名稱")

            for prog_name, group in grouped:
                first_row = group.iloc[0]
                times = sorted(list(set(group["演出場次時間"].tolist())))
                venues = list(dict.fromkeys([v for v in group["演出地點"].tolist() if v and v != "詳見售票頁"]))
                venues_str = "、".join(venues) if venues else first_row["演出地點"]

                m_d = re.search(r"(\d{4}[./-]\d{1,2}[./-]\d{1,2})", times[0]) if times else None
                first_play_date = m_d.group(1).replace(".", "-").replace("/", "-") if m_d else "即期"
                date_range_str = f"{times[0][:10]} ~ {times[-1][:10]}" if len(times) > 1 else (times[0] if times else "詳見售票頁")
                total_shows = len(group)

                summary_list.append({
                    "售票通路": first_row["售票通路"],
                    "所屬藝術節/系列": first_row["所屬藝術節/系列"],
                    "節目分類": first_row["節目分類"],
                    "演出單位": first_row["演出單位"],
                    "節目名稱": prog_name,
                    "首演日期": first_play_date,
                    "演出日期區間": f"{date_range_str} (共 {total_shows} 場)" if total_shows > 1 else date_range_str,
                    "巡迴/演出場館": venues_str,
                    "售票狀況/主要連結": first_row["售票狀況/連結"]
                })

            df_summary = pd.DataFrame(summary_list)

            # 核心排序演算法
            fest_first_dates = {}
            for fest, group in df_summary.groupby("所屬藝術節/系列"):
                fest_first_dates[fest] = group["首演日期"].min()

            df_summary["_fest_sort_date"] = df_summary.apply(
                lambda r: r["首演日期"] if r["所屬藝術節/系列"] == "單檔演出 (一般檔期)" else fest_first_dates.get(r["所屬藝術節/系列"], r["首演日期"]),
                axis=1
            )
            df_summary.sort_values(
                by=["_fest_sort_date", "所屬藝術節/系列", "首演日期", "節目名稱"],
                ascending=[True, True, True, True],
                inplace=True
            )
            df_summary.drop(columns=["_fest_sort_date"], inplace=True)
            df_details.sort_values(by=["演出場次時間", "所屬藝術節/系列", "節目名稱"], inplace=True)

            pbar.progress(95)
            status_box.info("正在產出 Excel 報表...")

            output = io.BytesIO()
            wb = openpyxl.Workbook()

            ws1 = wb.active
            ws1.title = "演出節目總表(合併)"
            ws1.append(list(df_summary.columns))
            for r in df_summary.to_dict(orient="records"):
                ws1.append(list(r.values()))

            ws2 = wb.create_sheet(title="場次明細表(展開)")
            ws2.append(list(df_details.columns))
            for r in df_details.to_dict(orient="records"):
                ws2.append(list(r.values()))

            def format_sheet(ws):
                header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
                header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
                center_align = Alignment(horizontal="center", vertical="center")
                left_align = Alignment(horizontal="left", vertical="center")
                border = Border(
                    left=Side(style='thin', color='E5E7EB'),
                    right=Side(style='thin', color='E5E7EB'),
                    top=Side(style='thin', color='E5E7EB'),
                    bottom=Side(style='thin', color='E5E7EB')
                )
                for col_idx in range(1, ws.max_column + 1):
                    c = ws.cell(row=1, column=col_idx)
                    c.fill = header_fill
                    c.font = header_font
                    c.alignment = center_align

                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        cell.border = border
                        cell.font = Font(name="微軟正黑體", size=10)
                        col_name = ws.cell(row=1, column=cell.column).value
                        if col_name in ["售票通路", "所屬藝術節/系列", "節目分類", "首演日期", "演出場次時間", "演出日期區間"]:
                            cell.alignment = center_align
                        else:
                            cell.alignment = left_align

                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = max(len(str(c.value or '')) for c in col)
                    ws.column_dimensions[col_letter].width = max(min(max_len * 1.6, 50), 14)

            format_sheet(ws1)
            format_sheet(ws2)

            wb.save(output)
            excel_bytes = output.getvalue()

            pbar.progress(100)
            pbar.empty()
            status_box.empty()

            st.success(f"🎉 統計完成！共自【{selected_platform_label}】擷取 **{len(df_summary)}** 檔演出節目、**{len(df_details)}** 個場次紀錄。")

            st.download_button(
                label=f"📥 下載【{selected_platform_label}】統計 Excel 報表",
                data=excel_bytes,
                file_name=f"表演藝術演出統計_{target_platform}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

            tab1, tab2 = st.tabs(["📊 演出節目總表 (去重合併)", "📝 場次明細表 (展開)"])
            with tab1:
                st.dataframe(df_summary, use_container_width=True, hide_index=True)
            with tab2:
                st.dataframe(df_details, use_container_width=True, hide_index=True)
